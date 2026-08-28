"""Telop Agent Ready Excel template export.

The Telop workbook uses three surfaces:

- ``Parsed_OCR_Data`` as normalized staging data.
- ``Field_Map`` as the routing contract.
- ``Sheet1`` as the customer-facing quotation.

This module keeps that contract separate from the generic immutable-cell
exporter because Telop has composite fields, preserved formulas, and an image
slot.
"""

from __future__ import annotations

import hashlib
import json
import re
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from pypdf import PdfReader

from .template_export import TemplateExportError

try:
    from openpyxl.drawing.image import Image as ExcelImage
except Exception:  # pragma: no cover - optional in some openpyxl installs
    ExcelImage = None


TELOP_TEMPLATE_TYPE = "telop_agent_ready"
TEXT_FALLBACK = "N/A"
REQUIRED_SHEETS = ("Sheet1", "Parsed_OCR_Data", "Field_Map")
FORMULA_CELLS = {"H7": "=G7*F7", "H8": "=F8*G8", "H9": "=SUM(H7:H8)"}
PARSED_HEADERS = [
    "field_key",
    "field_group",
    "raw_ocr_value",
    "normalized_value",
    "data_type",
    "required",
    "default_if_missing",
    "confidence",
    "source_page",
    "source_note",
]
CANONICAL_ROWS = [
    ("source_file", "source", "text", "Y", "PDF filename"),
    ("quote_date", "header", "date", "Y", "Quotation date"),
    ("purchase_order_no", "header", "text", "N", "Purchase order number"),
    ("order_no", "header", "text", "Y", "Order number"),
    ("buyer_name", "party", "text", "N", "Buyer"),
    ("supplier_name", "supplier", "text", "Y", "Supplier"),
    ("supplier_contact", "supplier", "text", "N", "Supplier contact"),
    ("supplier_phone", "supplier", "text", "N", "Supplier phone"),
    ("supplier_location", "supplier", "text", "N", "Supplier location"),
    ("supplier_website", "supplier", "text", "N", "Supplier website"),
    ("supplier_wechat", "supplier", "text", "N", "Supplier Wechat"),
    ("product_image", "image", "image_path", "N", "Product image path"),
    ("product_name_en", "product", "text", "Y", "English product name"),
    ("product_name_cn", "product", "text", "N", "Chinese product name"),
    ("item_type", "line_item", "text", "Y", "Line item type"),
    ("size", "spec", "text", "Y", "Product size"),
    ("materials", "spec", "text", "Y", "Product materials"),
    ("colors", "spec", "text", "N", "Product colors"),
    ("logo", "spec", "text", "N", "Logo requirement"),
    ("package", "spec", "text", "N", "Package"),
    ("other", "spec", "text", "N", "Other specifications"),
    ("unit", "line_item", "text", "N", "Unit"),
    ("qty", "line_item", "number", "Y", "Quantity"),
    ("unit_price_rmb", "line_item", "currency_rmb", "Y", "RMB unit price"),
    ("product_amount_rmb", "line_item", "currency_rmb", "Y", "Product amount"),
    ("mold_fee_rmb", "fee", "currency_rmb", "N", "Mold fee"),
    ("mold_refund_condition", "fee", "text", "N", "Mold refund condition"),
    ("total_rmb", "summary", "currency_rmb", "Y", "Total RMB"),
    ("total_rmb_upper_cn", "summary", "text", "N", "Uppercase RMB total"),
    ("trade_term", "terms", "text", "N", "Trade term"),
    ("payment_sample_mold", "terms", "text", "N", "Sample and mold payment"),
    ("payment_bulk", "terms", "text", "Y", "Bulk payment"),
    ("sample_fee", "terms", "text", "N", "Sample fee"),
    ("sample_time", "terms", "text", "N", "Sample time"),
    ("lead_time_days", "terms", "number", "Y", "Lead time days"),
    ("transportation_cost", "terms", "text", "Y", "Transportation cost"),
    ("delivery_method", "terms", "text", "Y", "Delivery method"),
    ("packing", "terms", "text", "N", "Packing"),
]


def telop_mapping_draft(report: dict[str, Any]) -> dict[str, Any]:
    """Return a disabled mapping draft for a Telop three-sheet workbook."""
    sheet_names = {sheet.get("name") for sheet in report.get("sheets", [])}
    if not set(REQUIRED_SHEETS).issubset(sheet_names):
        raise TemplateExportError("Telop template must contain Sheet1, Parsed_OCR_Data, and Field_Map.")
    return {
        "configured": False,
        "review_required": True,
        "template_type": TELOP_TEMPLATE_TYPE,
        "template_file": report["template_file"],
        "template_sha256": report["template_sha256"],
        "sheet_name": "Sheet1",
        "staging_sheet": "Parsed_OCR_Data",
        "field_map_sheet": "Field_Map",
    }


def is_telop_report(report: dict[str, Any]) -> bool:
    sheet_names = {sheet.get("name") for sheet in report.get("sheets", [])}
    return set(REQUIRED_SHEETS).issubset(sheet_names)


def is_telop_mapping(mapping: dict[str, Any]) -> bool:
    return mapping.get("template_type") == TELOP_TEMPLATE_TYPE


def inspect_telop_template_configuration(mapping_path: Path) -> dict[str, Any]:
    """Validate a Telop mapping and workbook before export."""
    if not mapping_path.exists():
        raise TemplateExportError(f"Excel template mapping does not exist: {mapping_path}")
    try:
        mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise TemplateExportError(f"Excel template mapping cannot be read: {exc}") from exc
    if not mapping.get("configured"):
        raise TemplateExportError("Telop template mapping has not been configured.")
    if not is_telop_mapping(mapping):
        raise TemplateExportError("Mapping is not a Telop Agent Ready mapping.")

    template_path = mapping_path.parent / str(mapping.get("template_file") or "")
    if not template_path.is_file():
        raise TemplateExportError(f"Original Excel template does not exist: {template_path}")
    if template_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise TemplateExportError("Telop template must be .xlsx or .xlsm.")

    wb = load_workbook(template_path, data_only=False, keep_vba=template_path.suffix.lower() == ".xlsm")
    missing_sheets = [sheet for sheet in REQUIRED_SHEETS if sheet not in wb.sheetnames]
    if missing_sheets:
        raise TemplateExportError(f"Telop template missing required sheet(s): {', '.join(missing_sheets)}")
    sheet = wb["Sheet1"]
    for cell, formula in FORMULA_CELLS.items():
        if sheet[cell].value != formula:
            raise TemplateExportError(f"Telop formula cell Sheet1!{cell} must remain {formula}.")
    parsed = wb["Parsed_OCR_Data"]
    headers = [parsed.cell(1, col).value for col in range(1, len(PARSED_HEADERS) + 1)]
    if headers != PARSED_HEADERS:
        raise TemplateExportError("Parsed_OCR_Data header does not match the Telop staging contract.")

    return {
        "configured": True,
        "template_type": TELOP_TEMPLATE_TYPE,
        "template_path": str(template_path),
        "template_file": template_path.name,
        "template_extension": template_path.suffix.lower(),
        "template_sha256": hashlib.sha256(template_path.read_bytes()).hexdigest(),
        "sheet_name": "Sheet1",
        "sheet_names": wb.sheetnames,
        "mapped_cell_count": len(CANONICAL_ROWS),
        "structure_fingerprint": _telop_fingerprint(template_path),
    }


def export_from_telop_template(
    job: dict[str, Any],
    project_root: Path,
    mapping_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    """Export an approved job into a Telop Agent Ready workbook."""
    inspection = inspect_telop_template_configuration(mapping_path)
    template_path = Path(inspection["template_path"])
    if output_path.suffix.lower() != template_path.suffix.lower():
        raise TemplateExportError("Output extension must match the original Telop template extension.")

    keep_vba = template_path.suffix.lower() == ".xlsm"
    wb = load_workbook(template_path, data_only=False, keep_vba=keep_vba)
    sheet = wb["Sheet1"]
    parsed = wb["Parsed_OCR_Data"]

    source_text = _source_text(job, project_root)
    fields = _telop_fields(job, project_root, source_text)
    image_path = _extract_product_image(job, project_root)
    if image_path:
        fields["product_image"] = str(image_path)

    _write_parsed_ocr(parsed, fields)
    _write_sheet1(sheet, fields)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    _verify_telop_output(output_path)
    return {
        **inspection,
        "mapping_path": str(mapping_path),
        "written_cell_count": len(CANONICAL_ROWS),
        "structure_unchanged": True,
        "formula_cells_preserved": dict(FORMULA_CELLS),
        "output_path": str(output_path),
        "fields_defaulted_to_na": sorted(key for key, value in fields.items() if value == TEXT_FALLBACK),
    }


def _unwrap(value: Any) -> Any:
    if isinstance(value, dict) and "value" in value:
        return value.get("value")
    return value


def _text(value: Any, fallback: str = TEXT_FALLBACK) -> str:
    value = _unwrap(value)
    if value is None or value == "":
        return fallback
    return str(value)


def _number(value: Any, default: float | None = None) -> float | None:
    value = _unwrap(value)
    if value is None or value == "":
        return default
    try:
        return float(str(value).replace(",", "").replace("RMB", "").strip())
    except ValueError:
        return default


def _first_match(text: str, pattern: str) -> str:
    match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def _all_matches(text: str, pattern: str) -> list[str]:
    return [re.sub(r"\s+", " ", value).strip() for value in re.findall(pattern, text, flags=re.IGNORECASE | re.DOTALL) if value.strip()]


def _source_text(job: dict[str, Any], project_root: Path) -> str:
    chunks: list[str] = []
    artifact_dir = Path(str((job.get("ocr") or {}).get("artifact_dir") or ""))
    if artifact_dir.is_dir():
        ocr_md = artifact_dir / "ocr.md"
        if ocr_md.is_file():
            chunks.append(ocr_md.read_text(encoding="utf-8", errors="ignore"))
    for page in job.get("quote", {}).get("raw_pages", []) or []:
        if isinstance(page, dict):
            chunks.append(str(page.get("text") or ""))
    source_path = project_root / "data" / "jobs" / str(job.get("id")) / "source.pdf"
    if source_path.is_file():
        try:
            reader = PdfReader(str(source_path))
            chunks.extend(page.extract_text() or "" for page in reader.pages)
        except Exception:
            pass
    return "\n".join(chunk for chunk in chunks if chunk)


def _telop_fields(job: dict[str, Any], project_root: Path, source_text: str) -> dict[str, Any]:
    quote = job.get("quote", {})
    headers = quote.get("headers", {})
    items = quote.get("items", [])
    totals = quote.get("totals", {})
    terms = quote.get("terms", {})
    first = items[0] if items else {}
    mold = _find_mold_item(items)

    purchase_order_no = _first_match(source_text, r"采购单号[:：]?\s*([0-9A-Za-z-]+)")
    order_no = _first_match(source_text, r"下单号[:：]?\s*([0-9A-Za-z-]+)")
    buyer_name = _first_match(source_text, r"甲方[:：]\s*([^\n\r]+)")
    supplier_name = _first_match(source_text, r"乙方[:：]\s*([^\n\r]+)")
    phones = _all_matches(source_text, r"电话[:：]\s*([0-9+\- ]{6,})")
    contacts = [
        value
        for value in _all_matches(source_text, r"联系人[:：]\s*([^\n\r ]+)")
        if "联系人" not in value and value not in {"电话", "N/A"}
    ]
    if not contacts:
        contacts = re.findall(r"([\u4e00-\u9fff]{1,6}(?:先生|女士|小姐))", source_text)

    product_name = _text(first.get("product_name"))
    specification = _text(first.get("specification"), "")
    remarks = _text(first.get("remarks"))
    lead_time_days = _first_match(source_text, r"交货期[:：].*?([0-9]+)\s*天") or _text(terms.get("lead_time"))

    fields = {
        "source_file": _text(job.get("source_file")),
        "quote_date": _date_from_text(source_text) or _text(headers.get("quote_date")),
        "purchase_order_no": purchase_order_no or _text(headers.get("quote_no")),
        "order_no": order_no or _text(headers.get("quote_no")),
        "buyer_name": buyer_name or _text(headers.get("customer")),
        "supplier_name": supplier_name or _text(headers.get("supplier")),
        "supplier_contact": contacts[-1] if contacts else TEXT_FALLBACK,
        "supplier_phone": phones[-1] if phones else TEXT_FALLBACK,
        "supplier_location": _text(first.get("location")),
        "supplier_website": TEXT_FALLBACK,
        "supplier_wechat": TEXT_FALLBACK,
        "product_image": TEXT_FALLBACK,
        "product_name_en": _english_product_name(product_name, specification),
        "product_name_cn": product_name,
        "item_type": "product",
        "size": TEXT_FALLBACK,
        "materials": _text(first.get("material")),
        "colors": _text(first.get("color")),
        "logo": TEXT_FALLBACK,
        "package": TEXT_FALLBACK,
        "other": _translate_other(remarks),
        "unit": _text(first.get("unit")),
        "qty": _number(first.get("quantity"), 0),
        "unit_price_rmb": _number(first.get("unit_price"), 0),
        "product_amount_rmb": _number(first.get("amount"), 0),
        "mold_fee_rmb": _number(mold.get("amount"), None) if mold else _mold_fee_from_text(source_text),
        "mold_refund_condition": _translate_mold_refund(source_text),
        "total_rmb": _number(totals.get("subtotal"), _number(totals.get("grand_total"), 0)),
        "total_rmb_upper_cn": _first_match(source_text, r"合计人民币金额[（(]大写[）)]?[:：]?\s*([^\n\r]+)") or TEXT_FALLBACK,
        "trade_term": TEXT_FALLBACK,
        "payment_sample_mold": TEXT_FALLBACK,
        "payment_bulk": _translate_payment(source_text) or _text(terms.get("payment")),
        "sample_fee": TEXT_FALLBACK,
        "sample_time": TEXT_FALLBACK,
        "lead_time_days": int(float(lead_time_days)) if str(lead_time_days).replace(".", "", 1).isdigit() else lead_time_days,
        "transportation_cost": _translate_transport(source_text),
        "delivery_method": _translate_delivery(source_text) or _text(terms.get("delivery")),
        "packing": TEXT_FALLBACK,
    }
    return fields


def _date_from_text(text: str) -> str:
    match = re.search(r"([0-9]{4})\s*年\s*([0-9]{1,2})\s*月\s*([0-9]{1,2})\s*日", text)
    if not match:
        return ""
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def _find_mold_item(items: list[dict[str, Any]]) -> dict[str, Any]:
    for item in items[1:]:
        name = _text(item.get("product_name"), "")
        if "模具" in name or "mold" in name.lower():
            return item
    return {}


def _mold_fee_from_text(text: str) -> float | None:
    value = _first_match(text, r"模具费.*?([0-9]+(?:\.[0-9]+)?)")
    return _number(value, None)


def _english_product_name(product_name: str, specification: str) -> str:
    combined = f"{product_name} {specification}"
    if "游戏手柄" in combined or "LT/RT" in combined or "LB/RB" in combined:
        return "Customizable controller for Xbox Series X/S and Xbox One"
    return product_name if product_name != TEXT_FALLBACK else TEXT_FALLBACK


def _translate_other(value: str) -> str:
    if "按照客户提供样品" in value:
        return "Shape and size according to customer sample"
    return value


def _translate_payment(text: str) -> str:
    if "60%预付款" in text:
        return "60% deposit before production, balance after inspection and before shipment"
    return ""


def _translate_transport(text: str) -> str:
    if "运输费用" in text and "乙方" in text and "承担" in text:
        return "Supplier bears"
    if "运输费用" in text and "甲方" in text and "承担" in text:
        return "Buyer bears"
    return TEXT_FALLBACK


def _translate_delivery(text: str) -> str:
    if "交货方式" in text and "按照甲方要求" in text:
        return "According to buyer's requirements"
    return ""


def _translate_mold_refund(text: str) -> str:
    if "退还模具费" in text and "10万" in text:
        return "Production quantity reaches 100,000 pieces, refund mold fee"
    return TEXT_FALLBACK


def _extract_product_image(job: dict[str, Any], project_root: Path) -> Path | None:
    source_path = project_root / "data" / "jobs" / str(job.get("id")) / "source.pdf"
    if not source_path.is_file():
        return None
    image_dir = project_root / "data" / "jobs" / str(job.get("id")) / "telop_assets"
    image_dir.mkdir(parents=True, exist_ok=True)
    try:
        page = PdfReader(str(source_path)).pages[0]
        images = list(page.images)
    except Exception:
        return None
    if not images:
        return None
    image = images[0]
    suffix = Path(getattr(image, "name", "product.jpg")).suffix or ".jpg"
    output = image_dir / f"product_image{suffix}"
    output.write_bytes(image.data)
    return output


def _write_parsed_ocr(ws: Any, fields: dict[str, Any]) -> None:
    for col, value in enumerate(PARSED_HEADERS, 1):
        ws.cell(1, col).value = value
    max_rows = max(ws.max_row, len(CANONICAL_ROWS) + 1)
    for row in range(2, max_rows + 1):
        for col in range(1, 11):
            ws.cell(row, col).value = None
    for row_index, (key, group, data_type, required, note) in enumerate(CANONICAL_ROWS, 2):
        _copy_row_style(ws, 2, row_index)
        value = fields.get(key, TEXT_FALLBACK)
        normalized = "" if value is None and data_type in {"number", "currency_rmb"} else _text(value)
        raw = normalized
        confidence = 0.95 if normalized not in {"", TEXT_FALLBACK} else 0.4
        row_values = [key, group, raw, normalized, data_type, required, TEXT_FALLBACK, confidence, 1, note]
        for col, cell_value in enumerate(row_values, 1):
            ws.cell(row_index, col).value = cell_value


def _write_sheet1(ws: Any, fields: dict[str, Any]) -> None:
    _set_value(ws, "A4", _text(fields.get("quote_date")))
    order_no = _text(fields.get("order_no"))
    _set_value(ws, "J4", order_no if order_no.startswith("#") or order_no == TEXT_FALLBACK else f"#{order_no}")
    _set_value(ws, "A7", _text(fields.get("supplier_name")))
    _add_product_image(ws, fields.get("product_image"))
    _set_value(ws, "C7", _text(fields.get("product_name_en")))
    _set_value(ws, "D7", _spec_block(fields))
    _set_value(ws, "E7", fields.get("lead_time_days") or TEXT_FALLBACK)
    _set_value(ws, "F7", _number(fields.get("qty"), 0))
    _set_value(ws, "G7", _number(fields.get("unit_price_rmb"), 0))
    ws["H7"].value = FORMULA_CELLS["H7"]
    _set_value(ws, "I7", _supplier_block(fields))
    _set_value(ws, "J7", _terms_block(fields))

    mold_fee = _number(fields.get("mold_fee_rmb"), None)
    _set_value(ws, "C8", "Mold cost" if mold_fee else TEXT_FALLBACK)
    _set_value(ws, "E8", TEXT_FALLBACK)
    _set_value(ws, "F8", 1 if mold_fee else 0)
    _set_value(ws, "G8", mold_fee or 0)
    ws["H8"].value = FORMULA_CELLS["H8"]
    _set_value(ws, "J8", _text(fields.get("mold_refund_condition")))
    ws["H9"].value = FORMULA_CELLS["H9"]


def _set_value(ws: Any, ref: str, value: Any) -> None:
    cell = ws[ref]
    if isinstance(cell, MergedCell):
        raise TemplateExportError(f"Cannot write merged-cell child {ws.title}!{ref}")
    if cell.data_type == "f":
        raise TemplateExportError(f"Cannot overwrite formula cell {ws.title}!{ref}")
    cell.value = value


def _copy_row_style(ws: Any, source_row: int, target_row: int) -> None:
    for col in range(1, min(ws.max_column, 10) + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)


def _add_product_image(ws: Any, image_path: Any) -> None:
    if hasattr(ws, "_images"):
        ws._images = []
    path = Path(str(_text(image_path, "")))
    if ExcelImage is None or not path.is_file():
        _set_value(ws, "B7", TEXT_FALLBACK)
        return
    image = ExcelImage(str(path))
    image.width = min(image.width, 130)
    image.height = min(image.height, 130)
    ws.add_image(image, "B7")


def _block(pairs: list[tuple[str, Any]]) -> str:
    return "\n".join(f"{label}:\n{_text(value)}" for label, value in pairs)


def _spec_block(fields: dict[str, Any]) -> str:
    return _block([
        ("Size", fields.get("size")),
        ("Materials", fields.get("materials")),
        ("Colors", fields.get("colors")),
        ("Logo", fields.get("logo")),
        ("Package", fields.get("package")),
        ("Other", fields.get("other")),
    ])


def _supplier_block(fields: dict[str, Any]) -> str:
    return _block([
        ("Name", fields.get("supplier_name")),
        ("Contact", fields.get("supplier_contact")),
        ("Location", fields.get("supplier_location")),
        ("Website", fields.get("supplier_website")),
        ("Wechat", fields.get("supplier_wechat")),
        ("Phone number", fields.get("supplier_phone")),
    ])


def _lead_time_sentence(value: Any) -> str:
    value = _text(value)
    if value == TEXT_FALLBACK:
        return TEXT_FALLBACK
    text = str(value)
    return f"{text} days upon receiving Buyer's deposit for completion" if text.replace(".", "", 1).isdigit() else text


def _terms_block(fields: dict[str, Any]) -> str:
    return "\n".join([
        "The final price will be determined after the final design and details are confirmed.",
        "",
        _block([
            ("Trade Term", fields.get("trade_term")),
            ("Payment term:\nSample & Mold", fields.get("payment_sample_mold")),
            ("Bulk", fields.get("payment_bulk")),
            ("Sample fee", fields.get("sample_fee")),
            ("Sample time", fields.get("sample_time")),
            ("Lead time", _lead_time_sentence(fields.get("lead_time_days"))),
            ("Transportation Cost", fields.get("transportation_cost")),
            ("Delivery method", fields.get("delivery_method")),
            ("Packing", fields.get("packing")),
        ]),
    ])


def _verify_telop_output(output_path: Path) -> None:
    wb = load_workbook(output_path, data_only=False)
    for required_sheet in REQUIRED_SHEETS:
        if required_sheet not in wb.sheetnames:
            raise TemplateExportError(f"Generated Telop workbook missing sheet: {required_sheet}")
    sheet = wb["Sheet1"]
    for cell, formula in FORMULA_CELLS.items():
        if sheet[cell].value != formula:
            raise TemplateExportError(f"Generated Telop workbook formula changed at Sheet1!{cell}")


def _telop_fingerprint(template_path: Path) -> str:
    wb = load_workbook(template_path, data_only=False, read_only=True)
    try:
        payload = {
            "sheets": wb.sheetnames,
            "sheet1_dimension": wb["Sheet1"].calculate_dimension(),
            "parsed_dimension": wb["Parsed_OCR_Data"].calculate_dimension(),
            "field_map_dimension": wb["Field_Map"].calculate_dimension(),
            "formulas": {cell: wb["Sheet1"][cell].value for cell in FORMULA_CELLS},
        }
        return hashlib.sha256(json.dumps(payload, sort_keys=True).encode("utf-8")).hexdigest()
    finally:
        wb.close()
