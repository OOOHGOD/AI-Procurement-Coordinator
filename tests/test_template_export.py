from __future__ import annotations

import json
import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from quote_assistant.service import QuoteService
from quote_assistant.template_export import MAIN_NS, TemplateExportError, export_from_immutable_template, inspect_template_configuration
from quote_assistant.template_inspect import inspect_xlsx_template, mapping_draft_from_report


WORKBOOK_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="报价单" sheetId="1" r:id="rId1"/></sheets>
</workbook>"""

RELS_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>"""

SHEET_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <cols><col min="1" max="3" width="18" customWidth="1"/></cols>
  <sheetData>
    <row r="2" ht="24" customHeight="1"><c r="A2" s="1" t="inlineStr"><is><t>报价编号</t></is></c><c r="B2" s="1" t="inlineStr"><is><t>原编号</t></is></c></row>
    <row r="4"><c r="A4" s="2" t="inlineStr"><is><t>序号</t></is></c><c r="B4" s="3" t="inlineStr"><is><t>产品名称</t></is></c></row>
    <row r="5"><c r="A5" s="2"><v>0</v></c><c r="B5" s="3" t="inlineStr"><is><t>预留</t></is></c></row>
    <row r="6"><c r="A6" s="2"><v>0</v></c><c r="B6" s="3" t="inlineStr"><is><t>预留</t></is></c></row>
    <row r="8"><c r="A8" s="4" t="inlineStr"><is><t>小计</t></is></c><c r="B8" s="4"><v>0</v></c><c r="C8" s="4"><f>SUM(B5:B6)</f><v>0</v></c></row>
  </sheetData>
  <mergeCells count="1"><mergeCell ref="D1:E1"/></mergeCells>
  <pageMargins left="0.2" right="0.2" top="0.5" bottom="0.5" header="0.2" footer="0.2"/>
  <pageSetup orientation="landscape" fitToWidth="1" fitToHeight="0"/>
</worksheet>"""


def create_template(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", WORKBOOK_XML)
        archive.writestr("xl/_rels/workbook.xml.rels", RELS_XML)
        archive.writestr("xl/worksheets/sheet1.xml", SHEET_XML)
        archive.writestr("xl/styles.xml", b"<styleSheet xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\"/>")
        archive.writestr("docProps/core.xml", b"<coreProperties/>")


def create_macro_template(path: Path) -> bytes:
    macro_bytes = b"macro-binary-placeholder"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", WORKBOOK_XML)
        archive.writestr("xl/_rels/workbook.xml.rels", RELS_XML)
        archive.writestr("xl/worksheets/sheet1.xml", SHEET_XML)
        archive.writestr("xl/styles.xml", b"<styleSheet xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\"/>")
        archive.writestr("xl/vbaProject.bin", macro_bytes)
        archive.writestr("docProps/core.xml", b"<coreProperties/>")
    return macro_bytes


def create_telop_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    workbook.create_sheet("Parsed_OCR_Data")
    workbook.create_sheet("Field_Map")
    for cell in ["A4", "J4", "A7", "B7", "C7", "D7", "E7", "F7", "G7", "I7", "J7", "C8", "E8", "F8", "G8", "J8"]:
        sheet[cell] = "placeholder"
    sheet["H7"] = "=G7*F7"
    sheet["H8"] = "=F8*G8"
    sheet["H9"] = "=SUM(H7:H8)"
    parsed = workbook["Parsed_OCR_Data"]
    parsed.append([
        "field_key",
        "field_group",
        "raw_ocr_value",
        "normalized_value",
        "data_type",
        "required",
        "default_if_missing",
        "confidence",
        "source_page",
        "source_note",
    ])
    field_map = workbook["Field_Map"]
    field_map.append(["map_key", "field_group", "source_field_keys", "target_sheet", "target_cell", "target_mode"])
    workbook.save(path)


class ImmutableTemplateTests(unittest.TestCase):
    def test_only_mapped_values_change(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "quote-template.xlsx"
            output = root / "output.xlsx"
            mapping_path = root / "mapping.json"
            create_template(template)
            mapping = {
                "configured": True,
                "template_file": template.name,
                "sheet_name": "报价单",
                "header_cells": {"quote.headers.quote_no.value": "B2"},
                "items": {"path": "quote.items", "start_row": 5, "max_rows": 2, "columns": {"line_no": "A", "product_name.value": "B"}},
                "total_cells": {"quote.totals.subtotal.value": "B8"},
                "clear_unused_item_rows": True,
            }
            mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
            job = {
                "quote": {
                    "headers": {"quote_no": {"value": "Q-001"}},
                    "items": [{"line_no": 1, "product_name": {"value": "Chair"}}],
                    "totals": {"subtotal": {"value": 1280.0}},
                }
            }
            audit = export_from_immutable_template(job, root, mapping_path, output)
            self.assertTrue(audit["structure_unchanged"])
            self.assertTrue(inspect_template_configuration(mapping_path)["configured"])

            with zipfile.ZipFile(template) as original, zipfile.ZipFile(output) as generated:
                self.assertEqual(set(original.namelist()), set(generated.namelist()))
                for name in original.namelist():
                    if name != "xl/worksheets/sheet1.xml":
                        self.assertEqual(original.read(name), generated.read(name))
                output_sheet = ET.fromstring(generated.read("xl/worksheets/sheet1.xml"))
            cells = {cell.get("r"): cell for cell in output_sheet.findall(f".//{{{MAIN_NS}}}c")}
            self.assertEqual("1", cells["A5"].find(f"{{{MAIN_NS}}}v").text)
            self.assertEqual("Chair", cells["B5"].find(f"{{{MAIN_NS}}}is/{{{MAIN_NS}}}t").text)
            self.assertIsNone(cells["B6"].find(f"{{{MAIN_NS}}}v"))
            self.assertEqual("SUM(B5:B6)", cells["C8"].find(f"{{{MAIN_NS}}}f").text)
            self.assertEqual("3", cells["B5"].get("s"))
            self.assertEqual("D1:E1", output_sheet.find(f".//{{{MAIN_NS}}}mergeCell").get("ref"))
            self.assertEqual("landscape", output_sheet.find(f".//{{{MAIN_NS}}}pageSetup").get("orientation"))

    def test_unconfigured_template_blocks_export(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            mapping_path = root / "mapping.json"
            mapping_path.write_text(json.dumps({"configured": False}), encoding="utf-8")
            with self.assertRaisesRegex(TemplateExportError, "尚未配置"):
                export_from_immutable_template({}, root, mapping_path, root / "output.xlsx")

    def test_more_items_than_reserved_rows_blocks_export(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "quote-template.xlsx"
            create_template(template)
            mapping_path = root / "mapping.json"
            mapping_path.write_text(json.dumps({
                "configured": True,
                "template_file": template.name,
                "sheet_name": "报价单",
                "header_cells": {"quote.headers.quote_no.value": "B2"},
                "items": {"path": "quote.items", "start_row": 5, "max_rows": 1, "columns": {"line_no": "A"}},
                "total_cells": {"quote.totals.subtotal.value": "B8"},
            }), encoding="utf-8")
            job = {"quote": {"headers": {"quote_no": {"value": "Q"}}, "items": [{"line_no": 1}, {"line_no": 2}], "totals": {"subtotal": {"value": 1}}}}
            with self.assertRaisesRegex(TemplateExportError, "禁止插行"):
                export_from_immutable_template(job, root, mapping_path, root / "output.xlsx")

    def test_xlsm_macro_template_is_preserved_without_changing_extension(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "quote-template.xlsm"
            macro_bytes = create_macro_template(template)
            mapping_path = root / "mapping.json"
            mapping = {
                "configured": True,
                "template_file": template.name,
                "sheet_name": "报价单",
                "header_cells": {"quote.headers.quote_no.value": "B2"},
                "items": {"path": "quote.items", "start_row": 5, "max_rows": 2, "columns": {"line_no": "A", "product_name.value": "B"}},
                "total_cells": {"quote.totals.subtotal.value": "B8"},
                "clear_unused_item_rows": True,
            }
            mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
            job = {
                "quote": {
                    "headers": {"quote_no": {"value": "Q-MACRO-001"}},
                    "items": [{"line_no": 1, "product_name": {"value": "Desk"}}],
                    "totals": {"subtotal": {"value": 2560.0}},
                }
            }

            output = root / "output.xlsm"
            audit = export_from_immutable_template(job, root, mapping_path, output)
            self.assertEqual(".xlsm", output.suffix.lower())
            self.assertTrue(audit["structure_unchanged"])
            with zipfile.ZipFile(output) as generated:
                self.assertEqual(macro_bytes, generated.read("xl/vbaProject.bin"))

            with self.assertRaisesRegex(TemplateExportError, "扩展名必须与原始模板完全一致"):
                export_from_immutable_template(job, root, mapping_path, root / "output.xlsx")

    def test_formula_and_duplicate_targets_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "quote-template.xlsx"
            create_template(template)
            mapping_path = root / "mapping.json"
            base = {
                "configured": True,
                "template_file": template.name,
                "sheet_name": "报价单",
                "header_cells": {"quote.headers.quote_no.value": "B2"},
                "items": {"path": "quote.items", "start_row": 5, "max_rows": 2, "columns": {"line_no": "A", "product_name.value": "B"}},
                "total_cells": {"quote.totals.subtotal.value": "C8"},
            }
            mapping_path.write_text(json.dumps(base), encoding="utf-8")
            with self.assertRaisesRegex(TemplateExportError, "公式单元格"):
                inspect_template_configuration(mapping_path)

            base["total_cells"] = {"quote.totals.subtotal.value": "B8"}
            base["header_cells"]["quote.headers.supplier.value"] = "B2"
            mapping_path.write_text(json.dumps(base), encoding="utf-8")
            with self.assertRaisesRegex(TemplateExportError, "多个字段映射"):
                inspect_template_configuration(mapping_path)

    def test_read_only_inspection_generates_disabled_draft(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "quote-template.xlsx"
            create_template(template)
            report = inspect_xlsx_template(template)
            draft = mapping_draft_from_report(report, "报价单")
            self.assertEqual(1, report["sheet_count"])
            self.assertEqual("B2", draft["header_cells"]["quote.headers.quote_no.value"])
            self.assertEqual("A", draft["items"]["columns"]["line_no"])
            self.assertEqual("B", draft["items"]["columns"]["product_name.value"])
            self.assertEqual(5, draft["items"]["start_row"])
            self.assertEqual("B8", draft["total_cells"]["quote.totals.subtotal.value"])
            self.assertFalse(draft["configured"])
            self.assertTrue(draft["review_required"])


class TelopTemplateTests(unittest.TestCase):
    def create_service(self, root: Path) -> QuoteService:
        (root / "templates").mkdir(parents=True)
        (root / "config.json").write_text(json.dumps({
            "confidence_threshold": 0.8,
            "critical_confidence_threshold": 0.9,
            "amount_tolerance": 0.02,
            "require_manual_approval": True,
            "alert_on_severities": ["error", "critical"],
            "critical_fields": ["supplier", "quote_no", "product_name", "quantity", "unit_price", "amount"],
            "excel_template_mapping": "templates/template_mapping.json",
        }), encoding="utf-8")
        (root / "templates" / "template_mapping.json").write_text(json.dumps({"configured": False}), encoding="utf-8")
        return QuoteService(root)

    def test_telop_template_import_activation_and_export(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            service = self.create_service(root)
            source_template = root / "telop.xlsx"
            create_telop_template(source_template)

            imported = service.import_template(source_template.name, source_template.read_bytes())
            draft = json.loads(Path(imported["draft_mapping_path"]).read_text(encoding="utf-8"))
            self.assertEqual("telop_agent_ready", draft["template_type"])
            self.assertFalse(draft["configured"])

            activation = service.activate_template_mapping({
                "reviewer": "Template Reviewer",
                "confirm_format_immutable": True,
                "mapping": draft,
            })
            self.assertEqual("telop_agent_ready", activation["template_type"])
            self.assertTrue(service.template_status()["configured"])

            job_id = "telopjob001"
            pdf_bytes = (Path(__file__).resolve().parents[1] / "samples" / "quote-normal.pdf").read_bytes()
            job_dir = root / "data" / "jobs" / job_id
            job_dir.mkdir(parents=True)
            (job_dir / "source.pdf").write_bytes(pdf_bytes)
            job = {
                "id": job_id,
                "revision": 1,
                "source_file": "input.pdf",
                "status": "approved",
                "source": {"sha256": hashlib.sha256(pdf_bytes).hexdigest(), "size_bytes": len(pdf_bytes)},
                "quote": {
                    "headers": {
                        "quote_no": {"value": "PO-1000"},
                        "supplier": {"value": "LDC Supplier"},
                        "customer": {"value": "Telop"},
                        "quote_date": {"value": "2026-06-22"},
                    },
                    "items": [{
                        "product_name": {"value": "Game controller"},
                        "quantity": {"value": 100},
                        "unit_price": {"value": 500},
                        "amount": {"value": 50000},
                        "remarks": {"value": "Custom sample"},
                    }, {
                        "product_name": {"value": "Mold fee"},
                        "amount": {"value": 14000},
                    }],
                    "totals": {"subtotal": {"value": 64000}},
                    "terms": {"payment": {"value": "60% deposit"}, "delivery": {"value": "Buyer requirement"}, "lead_time": {"value": 90}},
                    "raw_pages": [{"text": "采购单号：1000 下单号：123456\n甲方：Telop\n乙方：LDC Supplier\n电话：12345678901\n联系人：刘先生\n交货期：甲方支付预付款后90天完成\n运输费用：乙方公司承担\n交货方式：按照甲方要求\n生产数量达到10万个,退还模具费"}],
                },
                "validation": {"blocking_issue_count": 0, "warning_count": 0},
                "review": {"decision": "approved"},
                "review_history": [],
                "alerts": [],
                "alert": None,
                "export": None,
            }
            service.store.save(job)
            output = service.export_job(job_id)
            workbook = load_workbook(output, data_only=False)
            sheet = workbook["Sheet1"]
            self.assertEqual("2026-06-22", sheet["A4"].value)
            self.assertEqual("#123456", sheet["J4"].value)
            self.assertEqual("LDC Supplier", sheet["A7"].value)
            self.assertEqual(100, sheet["F7"].value)
            self.assertEqual(500, sheet["G7"].value)
            self.assertEqual("=G7*F7", sheet["H7"].value)
            self.assertEqual("=F8*G8", sheet["H8"].value)
            self.assertEqual("=SUM(H7:H8)", sheet["H9"].value)
            parsed = workbook["Parsed_OCR_Data"]
            self.assertEqual("source_file", parsed["A2"].value)
            self.assertEqual("input.pdf", parsed["D2"].value)


if __name__ == "__main__":
    unittest.main()
